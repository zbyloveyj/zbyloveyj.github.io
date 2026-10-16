#!/usr/bin/env python3
"""Create the full Chinese handbook as a printable PDF and editable Word file."""
from pathlib import Path
from html import escape
import json
import re
from bs4 import BeautifulSoup, NavigableString, Tag
import markdown
from weasyprint import HTML
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "downloads"
SOURCE = OUT / "complete-manuscript.md"
source = SOURCE.read_text(encoding="utf-8")
catalog = json.loads((ROOT / "data" / "catalog.json").read_text(encoding="utf-8"))

def validate():
    heads = re.findall(r"^# 第(?:\d+|[零一二三四五六七八九十]+)章",source,re.M)
    if len(heads) != 50: raise ValueError("Source must contain exactly 50 chapters")
    if len(catalog["volumes"]) != 9: raise ValueError("Missing book volumes")

PART_INTROS = [
  "认识研究工作的组织、证据与交流方式，是选择专业和理解科研训练的起点。本编从读研目标、学术生态、培养类型与方向选择出发，帮助读者形成有依据的个人判断。",
  "正式报考不应依赖往届截图或未经核实的信息。学习如何查找政策、比较培养项目、准备申请材料并与潜在导师进行专业沟通。",
  "入学第一年不必一次掌握所有复杂技术。理解培养要求，建立研究记录与安全意识，在文献阅读和小型实操中逐步形成可靠的工作习惯。",
  "真正的文献阅读既要获取知识，也要判断证据的来源与局限。本编训练检索、精读、知识管理、批判性评价和具有综合判断的综述写作。",
  "研究方案的价值是将值得回答的问题转化为可实施的证据设计。从科学假设、研究类型、样本与变量，一直推进到开题报告和风险管理。",
  "科研数据可能包含测量误差、缺失、批次差异与复杂依赖结构。通过实际操作学习实验质量、数据清洗、统计推断、机器学习和可重复绘图。",
  "论文写作不是把图表逐个拼接，而是使研究问题、方法、直接证据和结论相互对应。本编从全文结构进入摘要、Methods、Results、Discussion和同行评议。",
  "科研活动需要合作、诚信、明确责任和持续交流。理解如何正确使用AI工具，维护数据和作者权益，并将真实项目转化为可检查的成果。",
  "完成学位是一个阶段的结束，也是新的能力起点。本编讨论培养节点、学位论文、送审答辩、求职和长期发展，不以单一评价指标定义研究生。"
]

def html_body():
    txt = markdown.markdown(source, extensions=["tables","fenced_code"])
    soup = BeautifulSoup(txt, "html.parser")
    chap = 0
    parts = 0
    for h in soup.find_all("h1"):
        label = h.get_text(" ",strip=True)
        if label == "从读研到做研究":
            h.decompose()
            continue
        if re.match(r"^第[一二三四五六七八九]编",label):
            parts += 1
            h["class"] = "part-heading"
            h["id"] = "part-" + str(parts)
            panel = soup.new_tag("section")
            panel["class"] = "part-intro"
            note = soup.new_tag("p")
            note["class"] = "part-lead"
            note.string = PART_INTROS[parts-1]
            panel.append(note)
            chapter_label = soup.new_tag("p")
            chapter_label["class"] = "part-kicker"
            chapter_label.string = "本编章节 | CONTENTS"
            panel.append(chapter_label)
            listing = soup.new_tag("ul")
            listing["class"] = "part-chapters"
            for item in catalog["chapters"]:
                if item["part"] == parts:
                    entry = soup.new_tag("li")
                    entry.string = f'{item["n"]:02d}　{item["title"]}'
                    listing.append(entry)
            panel.append(listing)
            h.insert_after(panel)
        elif re.match(r"^第(?:\d+|[零一二三四五六七八九十]+)章",label):
            chap += 1
            h["class"] = "chapter-heading"
            h["id"] = "chapter-" + str(chap)
        elif "前言" in label:
            h["class"] = "front-heading"
    if chap != 50 or parts != 9:
        raise ValueError(f"HTML conversion lost chapters/parts: {chap}/{parts}")
    return soup

def build_html(soup):
    part_chapters = []
    names = "一二三四五六七八九"
    for i,title in enumerate(catalog["volumes"],1):
        part_chapters.append("<div class='toc-part'>第"+names[i-1]+"编　"+escape(title)+"</div>")
        for ch in catalog["chapters"]:
            if ch["part"]==i:
                part_chapters.append("<div class='toc-row'><span>"+str(ch["n"]).zfill(2)+"</span><a href='#chapter-"+str(ch["n"])+"'>"+escape(ch["title"])+"</a></div>")
    css = """
@page:first { @top-right { content: none; } }
@page {
  size: A5;
  margin: 18mm 16mm 18mm 16mm;
  @top-right {content:"从读研到做研究";font-size:7pt;color:#5c7368;}
  @bottom-center {content: counter(page);font-size:8pt;color:#5c7368;}
}
html {font-family:"Noto Serif CJK SC","Noto Serif CJK",serif;color:#25382f;}
body {font-size:10pt;line-height:1.78;}
p {text-align:justify;margin:0 0 3.5mm;orphans:3;widows:3;overflow-wrap:anywhere;}
.cover {page-break-after:always;height:150mm;display:flex;flex-direction:column;justify-content:space-between;}
.cover .label {font-family:sans-serif;letter-spacing:.18em;color:#557d6c;font-size:9pt;}
.cover h1 {font-size:31pt;line-height:1.4;color:#263e32;margin:45mm 0 0;}
.cover p {color:#617166;text-align:left;}
.toc {page-break-after:always;}
.toc h1 {font-size:22pt;color:#263e32;margin:0 0 14mm;}
.toc-part {font-weight:bold;color:#477461;margin-top:4mm;font-size:11pt;border-top:1px solid #d9e1d7;padding-top:3mm;}
.toc-row {font-size:9pt;margin:1.4mm 0 1.4mm 3mm;}
.toc-row span {display:inline-block;width:9mm;color:#7a8d81;}
a {color:#426b57;text-decoration:none;}
h1.part-heading {page-break-before:always;page-break-after:avoid;font-size:21pt;color:#47715e;margin-top:23mm;margin-bottom:7mm;line-height:1.5;}
.part-intro {page-break-inside:avoid;margin-top:6mm;}
.part-intro .part-lead {font-size:10pt;line-height:1.82;color:#51675a;text-align:left;margin:0 0 13mm;}
.part-intro .part-kicker {font:700 8pt sans-serif;letter-spacing:.13em;color:#759081;margin:0 0 3mm;}
.part-intro ul.part-chapters {list-style:none;padding:0;margin:0;font-size:9.5pt;line-height:1.85;}
.part-intro ul.part-chapters li {border-bottom:.4pt solid #e4eae1;padding:1.8mm 0;margin:0;}
h1.chapter-heading {page-break-before:always;page-break-after:avoid;font-size:17pt;line-height:1.55;color:#263e32;margin:0 0 12mm;}
h1.front-heading {page-break-before:always;font-size:19pt;color:#304e3e;}
h1:not(.part-heading):not(.chapter-heading):not(.front-heading) {font-size:16pt;page-break-before:always;color:#304e3e;}
h2 {font-size:12.2pt;color:#365e4a;margin:8mm 0 3mm;page-break-after:avoid;}
h3 {font-size:10.8pt;color:#2f4c3d;margin:5mm 0 2mm;page-break-after:avoid;}
strong {color:#263e32;}
blockquote {background:#f5f7f3;border-left:2px solid #63856f;padding:2mm 4mm;margin:5mm 0;color:#4c6153;}
ul,ol {padding-left:7mm;margin:3mm 0 5mm;}
li {margin:1.5mm 0;}
table {border-collapse:collapse;width:100%;font-size:8pt;margin:5mm 0;table-layout:fixed;}
th,td {border:.4pt solid #d5dfd6;padding:1.6mm;vertical-align:top;overflow-wrap:anywhere;}
th {background:#f0f4ee;}
pre {background:#f6f8f4;border:1px solid #e2e6df;padding:2.5mm;white-space:pre-wrap;overflow-wrap:anywhere;font-size:7.1pt;line-height:1.45;}
code {font-size:.9em;overflow-wrap:anywhere;}
hr {display:none;} /* Chapter boundary rules plus forced H1 page breaks created blank pages. */
"""
    return "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><style>"+css+"</style></head><body>"+(
      "<section class='cover'><div class='label'>A RESEARCHER'S FIELD GUIDE · 2026</div>"
      "<div><h1>从读研<br>到做研究</h1><p>研究生学术入门手册 · 九编五十章</p></div>"
      "<p>著 · 一心问道<br>内容深化编辑版 V2.1 · 2026年10月</p></section>"
      "<section class='toc'><h1>目录</h1>"+''.join(part_chapters)+"</section>"
      +str(soup)+"</body></html>")

def add_rich(paragraph,tag):
    for node in tag.children:
        if isinstance(node,NavigableString):
            paragraph.add_run(str(node))
        elif isinstance(node,Tag):
            if node.name in ("strong","b","em","i","code"):
                run=paragraph.add_run(node.get_text())
                if node.name in ("b","strong"):run.bold=True
                if node.name in ("em","i"):run.italic=True
            else:
                paragraph.add_run(node.get_text())

def set_font(style,pt,bold=False,color=None):
    style.font.name="Noto Serif CJK SC"
    style.font.size=Pt(pt)
    style.font.bold=bold
    if color: style.font.color.rgb=RGBColor(*color)
    rpr=style.element.get_or_add_rPr()
    fonts=rpr.rFonts
    if fonts is None:
        fonts=OxmlElement("w:rFonts")
        rpr.insert(0,fonts)
    fonts.set(qn("w:eastAsia"),"Noto Serif CJK SC")

def build_word(soup,target):
    d=Document()
    d.core_properties.title="从读研到做研究"
    d.core_properties.author="一心问道"
    sec=d.sections[0]
    sec.page_width=Cm(14.8); sec.page_height=Cm(21)
    sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.6)
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.6)
    for name,pt,bold,color in [("Normal",10.5,False,None),("Title",27,True,(42,72,54)),("Heading 1",18,True,(42,72,54)),("Heading 2",13,True,(56,99,77)),("Heading 3",11,True,None)]:
        set_font(d.styles[name],pt,bold,color)
    d.styles["Normal"].paragraph_format.line_spacing=1.45
    d.styles["Normal"].paragraph_format.space_after=Pt(6)
    d.styles["Normal"].paragraph_format.first_line_indent=Pt(18)
    d.add_paragraph("\n\n\n\n")
    d.add_heading("从读研到做研究",0)
    p=d.add_paragraph("研究生学术入门手册\n九编五十章 · 内容深化编辑版 V2.1")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("\n\n\n\n")
    p=d.add_paragraph("一心问道\n2026年10月")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    d.add_heading("目录",1)
    chinese="一二三四五六七八九"
    for part,title in enumerate(catalog["volumes"],1):
        d.add_paragraph("第"+chinese[part-1]+"编　"+title,style="Heading 2")
        for ch in catalog["chapters"]:
            if ch["part"]==part:
                d.add_paragraph("第"+str(ch["n"]).zfill(2)+"章　"+ch["title"])
    d.add_page_break()
    seen=0
    for block in list(soup.children):
        if not isinstance(block,Tag):continue
        tag=block.name
        if tag=="h1":
            name=block.get_text(" ",strip=True)
            if "编" in name[:5] and name.startswith("第"):
                d.add_page_break();d.add_heading(name,1)
            elif re.match(r"^第(?:\d+|[零一二三四五六七八九十]+)章",name):
                seen+=1;d.add_page_break();d.add_heading(name,1)
            else:d.add_heading(name,1)
        elif tag=="section" and "part-intro" in block.get("class",[]):
            intro=block.find("p",class_="part-lead")
            if intro:
                p=d.add_paragraph(intro.get_text(" ",strip=True))
                p.style="Quote" if "Quote" in d.styles else "Normal"
            for li in block.find_all("li"):
                d.add_paragraph(li.get_text(" ",strip=True))
        elif tag in ("h2","h3"):
            d.add_heading(block.get_text(),2 if tag=="h2" else 3)
        elif tag=="p":
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;add_rich(p,block)
        elif tag=="blockquote":
            p=d.add_paragraph(block.get_text(" ",strip=True))
            p.style="Quote" if "Quote" in d.styles else "Normal"
        elif tag in ("ul","ol"):
            sty="List Bullet" if tag=="ul" else "List Number"
            for li in block.find_all("li",recursive=False):
                d.add_paragraph(li.get_text(" ",strip=True),style=sty)
        elif tag=="pre":
            p=d.add_paragraph(block.get_text())
            p.style="No Spacing"
            for run in p.runs:run.font.name="Consolas";run.font.size=Pt(8)
        elif tag=="table":
            rows=block.find_all("tr")
            if not rows:continue
            first=rows[0].find_all(["td","th"],recursive=False)
            if not first:continue
            t=d.add_table(rows=0,cols=len(first));t.style="Table Grid"
            for row in rows:
                cells=row.find_all(["td","th"],recursive=False)
                dest=t.add_row().cells
                for i,c in enumerate(cells[:len(dest)]):dest[i].text=c.get_text(" ",strip=True)
        elif tag=="hr":
            d.add_paragraph("")
    if seen!=50:raise ValueError(f"Word output has {seen} chapters")
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("— ")
    fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE")
    footer._p.append(fld)
    footer.add_run(" —")
    d.save(target)

def main():
    validate()
    soup=html_body()
    page=build_html(soup)
    pdf=OUT/"complete-book.pdf"
    word=OUT/"complete-book.docx"
    print("Rendering complete PDF...")
    HTML(string=page,base_url=str(ROOT)).write_pdf(str(pdf))
    print("Generating editable Word...")
    build_word(soup,word)
    from pypdf import PdfReader
    pages=len(PdfReader(str(pdf)).pages)
    # Detect virtually blank pages; previous editions contained empty pages
    # because an HR could spill to a new page just before forced chapter breaks.
    reader=PdfReader(str(pdf))
    nearly_empty=[]
    for i,page in enumerate(reader.pages):
        extracted=(page.extract_text() or "").strip()
        if len(extracted)<20:
            nearly_empty.append(i+1)
    if nearly_empty:
        raise ValueError(f"PDF contains blank/near-blank pages: {nearly_empty[:15]}")
    if pages<100 or pdf.stat().st_size<100_000 or word.stat().st_size<50_000:
        raise ValueError(f"Unexpected output sizes: {pages}, {pdf.stat().st_size}, {word.stat().st_size}")
    print(f"SUCCESS: {pages} PDF pages; PDF {pdf.stat().st_size:,} bytes; DOCX {word.stat().st_size:,} bytes")

if __name__=="__main__":
    main()
